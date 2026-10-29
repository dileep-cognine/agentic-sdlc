import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services.document_parser import DocumentError, parse_document


def test_txt_validation_and_normalization():
    text, sources = parse_document("input.txt", "text/plain", b"\xef\xbb\xbfHello\r\nWorld")
    assert text == "Hello\nWorld"
    assert sources[1]["location"] == "TXT, line 2"
    assert sources[1]["id"] == "S2"


def test_docx_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Portal login")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Customer"
    data = io.BytesIO()
    document.save(data)
    text, sources = parse_document(
        "input.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data.getvalue(),
    )
    assert "Portal login" in text and "Role | Customer" in text
    assert sources[0]["text"] == "Portal login"


def test_pdf_page_provenance():
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 700 Td (Portal login.) Tj ET")
    page[NameObject("/Contents")] = stream
    data = io.BytesIO()
    writer.write(data)
    text, sources = parse_document("input.pdf", "application/pdf", data.getvalue())
    assert text == "Portal login."
    assert sources[0]["location"] == "PDF page 1, line 1"


@pytest.mark.parametrize(
    "name,mime,content",
    [
        ("evil.exe", "text/plain", b"Hello"),
        ("fake.pdf", "application/pdf", b"Hello"),
        (
            "fake.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"PKbad",
        ),
        ("file.txt", "text/plain", b"\xff\x00"),
        ("file.txt", "text/plain", b"hello\x00world"),
        ("file.txt", "application/pdf", b"Hello"),
        ("empty.txt", "text/plain", b""),
        ("huge.txt", "text/plain", b"x" * (5 * 1024 * 1024 + 1)),
    ],
)
def test_reject_invalid_uploads(name, mime, content):
    with pytest.raises(DocumentError):
        parse_document(name, mime, content)


def test_docx_zip_expansion_limit():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (21 * 1024 * 1024))
    with pytest.raises(DocumentError, match="expanded size"):
        parse_document(
            "input.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data.getvalue(),
        )


def test_upload_api(client):
    api, _ = client
    project = api.post("/api/projects", json={"name": "Upload"}).json()
    path = f"/api/projects/{project['id']}/uploads"
    response = api.post(
        path,
        data={"expected_version": "1"},
        files={"file": ("input.txt", b"Portal login", "text/plain")},
    )
    assert response.status_code == 201, response.text
    assert response.json()["text"] == "Portal login"
    assert (
        api.post(
            path,
            data={"expected_version": "1"},
            files={"file": ("input.txt", b"stale", "text/plain")},
        ).status_code
        == 409
    )


def test_request_body_limit(client):
    api, _ = client
    response = api.post(
        "/api/projects",
        content=b"x" * (6 * 1024 * 1024 + 1),
        headers={"Content-Type": "application/json"},
    )
    assert response.status_code == 413
